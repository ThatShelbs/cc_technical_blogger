from __future__ import annotations

import io
from collections.abc import Iterator
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from proseferno.config import Settings
from proseferno.demo import FakeLLM
from proseferno.llm.auth import ClaudeAuth
from proseferno.main import create_app
from proseferno.secrets import MemorySecretStore
from proseferno.services import Services, build_services

FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    seed = tmp_path / "seed"
    seed.mkdir()
    return Settings(
        data_dir=tmp_path / "data",
        seed_samples_dir=seed,
        frontend_dist=tmp_path / "dist",
        _env_file=None,  # type: ignore[call-arg]
    )


@pytest.fixture
def fake_llm() -> FakeLLM:
    return FakeLLM()


@pytest.fixture
def services(settings: Settings, fake_llm: FakeLLM, tmp_path: Path) -> Services:
    secrets = MemorySecretStore()
    auth = ClaudeAuth(
        secrets,
        env={"APPDATA": str(tmp_path / "appdata"), "ANTHROPIC_CONFIG_DIR": str(tmp_path / "anthropic")},
        discover_ant=False,
    )
    return build_services(settings, secrets=secrets, llm=fake_llm, auth=auth, inline_jobs=True)


@pytest.fixture
def client(settings: Settings, services: Services) -> Iterator[TestClient]:
    with TestClient(create_app(settings, services)) as test_client:
        yield test_client


def medium_export_docx() -> bytes:
    """A small public-domain post laid out like a Medium export (no heading styles, UI chrome, captions)."""
    document = Document()
    for text in [
        "Pie Charts Are Lame",
        "A short roast of the most popular chart in the office.",
        "Press enter or click to view image in full size",
    ]:
        document.add_paragraph(text)
    document.add_paragraph().add_run().add_picture(io.BytesIO(_tiny_png()))
    document.add_paragraph("Source: Office survey; Author: Test Writer; Made with Tableau")
    document.add_paragraph(
        "Pie charts are everywhere. I see them in every quarterly deck, and I have opinions. "
        "Most of the time, a bar chart would tell the story faster. You know it, I know it."
    )
    document.add_paragraph("Why Pie Charts Fail")
    document.add_paragraph(
        "Humans are bad at comparing angles. We are much better at comparing lengths, which is why bars win. "
        "So you're telling me the biggest slice is the biggest? Shocking!"
    )
    document.add_paragraph("What is the biggest slice?")
    document.add_paragraph().add_run().add_picture(io.BytesIO(_tiny_png()))
    document.add_paragraph("A Better Option")
    document.add_paragraph(
        "Try a sorted bar chart. Label the bars directly, drop the legend, and let the reader breathe. "
        "Your audience will thank you, and your designer will stop sighing."
    )
    document.add_paragraph("42")
    document.add_paragraph("Follow")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _tiny_png() -> bytes:
    from proseferno.demo import png_bytes

    return png_bytes(16, 16)


@pytest.fixture
def sample_docx() -> bytes:
    return medium_export_docx()
