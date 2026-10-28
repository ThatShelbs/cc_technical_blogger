from __future__ import annotations

import io
from pathlib import Path

from pptx import Presentation

from proseferno.demo import DRAFT, png_bytes
from proseferno.export.linkedin import POST_LIMIT, carousel_pdf, carousel_pptx, carousel_slides, fit_post, hashtags, template_post
from proseferno.export.render import Figure
from proseferno.export.x import teaser_post


def figures(tmp_path: Path) -> dict[str, Figure]:
    out = {}
    for slot in ("hero", "img2", "img3", "meme"):
        path = tmp_path / f"{slot}.png"
        path.write_bytes(png_bytes(160, 90))
        out[slot] = Figure(slot, path.name, path, f"{slot} alt", f"{slot} caption", "Author: Me", 160, 90)
    return out


def test_template_post_fits_linkedin() -> None:
    post = template_post(DRAFT, ["Data Visualization", "Charts", "dataviz"])
    assert post.startswith("Pie charts have a fan club.")
    assert "→ Why People Love Pies" in post and "#DataVisualization #Charts #Dataviz" in post
    assert len(post) <= POST_LIMIT and "{{image" not in post and "##" not in post


def test_fit_post_trims_and_strips_markdown() -> None:
    long = "## Big **bold** idea\n\n" + "\n".join("A line of text that goes on." for _ in range(200))
    fitted = fit_post(long)
    assert len(fitted) <= POST_LIMIT and fitted.startswith("Big bold idea") and fitted.endswith("…")
    assert hashtags(["machine learning", "LLMs", "AI"]) == ["#MachineLearning", "#LLMs", "#AI"]


def test_carousel_pdf_and_pptx(tmp_path: Path) -> None:
    import pymupdf

    slides = carousel_slides(DRAFT, figures(tmp_path), "Shelby Temple")
    kinds = [s.kind for s in slides]
    assert kinds[0] == "cover" and kinds[-1] == "closing" and "section" in kinds and "image" in kinds
    assert slides[0].title == "Pie Charts Are Lame" and slides[-1].lines[1] == "by Shelby Temple"

    with pymupdf.open(stream=carousel_pdf(slides, "Pie Charts Are Lame"), filetype="pdf") as doc:
        assert doc.page_count == len(slides)
        assert (doc[0].rect.width, doc[0].rect.height) == (1080, 1350)
        assert "Pie Charts Are Lame" in doc[0].get_text()

    deck = Presentation(io.BytesIO(carousel_pptx(slides)))
    assert len(deck.slides) == len(slides)
    assert any("Why People Love Pies" in shape.text_frame.text for shape in deck.slides[1].shapes if shape.has_text_frame)


def test_x_teaser_fits_one_post() -> None:
    post = (
        "# Pie Charts Are Lame\n\n*A roast.*\n\n{{image:hero}}\n\n"
        + "Pie charts have a fan club and I am not in it. " * 20
        + "\n\n## Next\n\nMore."
    )
    text = teaser_post(post)
    assert len(text) <= 280 and text.endswith("“Pie Charts Are Lame”") and text.startswith("Pie charts have a fan club")
    assert teaser_post("# Only a title\n") == "New article: “Only a title”"
