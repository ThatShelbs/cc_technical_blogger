"""Parsers for every supported sample format: .docx, .pdf, .md, .txt, .html and URLs."""

from __future__ import annotations

import io
import re
from pathlib import Path

from docx import Document as load_docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from markdown_it import MarkdownIt

from proseferno.ingest.cleanup import clean_document
from proseferno.ingest.document import Block, ParsedDocument

SUPPORTED_EXTENSIONS = {".docx", ".pdf", ".md", ".markdown", ".txt", ".html", ".htm"}

_MONO_FONTS = ("courier", "consolas", "menlo", "monaco", "source code", "fira code", "mono")


class UnsupportedFormatError(ValueError):
    pass


def parse_file(name: str, data: bytes) -> ParsedDocument:
    """Parse an uploaded file and return a cleaned document."""
    ext = Path(name).suffix.lower()
    if ext == ".docx":
        doc = parse_docx(data)
    elif ext == ".pdf":
        doc = parse_pdf(data)
    elif ext in (".md", ".markdown", ".txt"):
        doc = parse_markdown(data.decode("utf-8", errors="replace"), source_type=ext.lstrip("."))
    elif ext in (".html", ".htm"):
        doc = parse_html(data.decode("utf-8", errors="replace"))
    else:
        raise UnsupportedFormatError(f"Unsupported file type '{ext}'. Use one of: {', '.join(sorted(SUPPORTED_EXTENSIONS))}.")
    return clean_document(doc)


# -- .docx -------------------------------------------------------------------


def parse_docx(data: bytes) -> ParsedDocument:
    document = load_docx(io.BytesIO(data))
    doc = ParsedDocument(source_type="docx")
    props = document.core_properties
    doc.author = (props.author or "").strip()
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            _docx_paragraph(Paragraph(child, document), doc)
        elif child.tag == qn("w:tbl"):
            for row in child.iter(qn("w:tr")):
                cells = ["".join(t.text or "" for t in cell.iter(qn("w:t"))).strip() for cell in row.iter(qn("w:tc"))]
                if any(cells):
                    doc.blocks.append(Block("paragraph", " | ".join(cells)))
    return doc


def _docx_paragraph(para: Paragraph, doc: ParsedDocument) -> None:
    element = para._element
    images = len(element.findall(".//" + qn("w:drawing"))) + len(element.findall(".//" + qn("w:pict")))
    text = para.text.strip()
    style = (para.style.name if para.style is not None else "") or ""
    style_l = style.lower()

    if text:
        if style_l == "title" and not doc.title:
            doc.title = text
        elif style_l == "subtitle" and not doc.subtitle:
            doc.subtitle = text
        elif style_l.startswith("heading"):
            level = int(m.group()) if (m := re.search(r"\d", style)) else 2
            doc.blocks.append(Block("heading", text, level=max(2, min(level + 1, 4))))
        elif "quote" in style_l:
            doc.blocks.append(Block("quote", text))
        elif "code" in style_l or "preformatted" in style_l or _is_monospace(para):
            doc.blocks.append(Block("code", para.text.rstrip()))
        elif "list" in style_l or element.find(".//" + qn("w:numPr")) is not None:
            doc.blocks.append(Block("list_item", text, ordered="number" in style_l))
        else:
            doc.blocks.append(Block("paragraph", text))
    for _ in range(images):
        doc.blocks.append(Block("image"))


def _is_monospace(para: Paragraph) -> bool:
    runs = [r for r in para.runs if r.text.strip()]
    if not runs:
        return False
    for run in runs:
        font = (run.font.name or "").lower()
        if not any(mono in font for mono in _MONO_FONTS):
            return False
    return True


# -- .pdf --------------------------------------------------------------------


def parse_pdf(data: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
    doc = ParsedDocument(source_type="pdf")
    meta = reader.metadata
    if meta is not None:
        doc.title = (meta.title or "").strip()
        doc.author = (meta.author or "").strip()
    for chunk in re.split(r"\n\s*\n", text):
        para = " ".join(line.strip() for line in chunk.splitlines() if line.strip())
        if para:
            doc.blocks.append(Block("paragraph", para))
    return doc


# -- markdown / text ---------------------------------------------------------


def parse_markdown(text: str, source_type: str = "md") -> ParsedDocument:
    md = MarkdownIt("commonmark")
    tokens = md.parse(text)
    doc = ParsedDocument(source_type=source_type)
    list_stack: list[bool] = []
    i = 0
    while i < len(tokens):
        tok = tokens[i]
        if tok.type in ("bullet_list_open", "ordered_list_open"):
            list_stack.append(tok.type == "ordered_list_open")
        elif tok.type in ("bullet_list_close", "ordered_list_close"):
            list_stack.pop()
        elif tok.type == "heading_open":
            level = int(tok.tag[1])
            content = tokens[i + 1].content.strip()
            if level == 1 and not doc.title:
                doc.title = content
            else:
                doc.blocks.append(Block("heading", content, level=max(2, level)))
            i += 2
        elif tok.type == "inline":
            content = tok.content.strip()
            if not content:
                pass
            elif re.fullmatch(r"!\[[^\]]*\]\([^)]*\)", content):
                doc.blocks.append(Block("image"))
                alt = re.match(r"!\[([^\]]*)\]", content)
                if alt and alt.group(1).strip():
                    doc.blocks.append(Block("caption", alt.group(1).strip()))
            elif list_stack:
                doc.blocks.append(Block("list_item", content, ordered=list_stack[-1]))
            elif tokens[i - 1].type == "paragraph_open" and _inside_blockquote(tokens, i):
                doc.blocks.append(Block("quote", content))
            else:
                doc.blocks.append(Block("paragraph", content))
        elif tok.type in ("fence", "code_block"):
            doc.blocks.append(Block("code", tok.content.rstrip("\n")))
        i += 1
    return doc


def _inside_blockquote(tokens: list, index: int) -> bool:  # type: ignore[type-arg]
    depth = 0
    for tok in tokens[:index]:
        if tok.type == "blockquote_open":
            depth += 1
        elif tok.type == "blockquote_close":
            depth -= 1
    return depth > 0


# -- html / url ----------------------------------------------------------------


def parse_html(html: str, url: str | None = None) -> ParsedDocument:
    import trafilatura

    extracted = trafilatura.extract(
        html,
        url=url,
        output_format="markdown",
        include_images=True,
        include_formatting=True,
        include_links=False,
        with_metadata=False,
    )
    if not extracted:
        raise ValueError("Could not find article text on that page.")
    doc = parse_markdown(extracted, source_type="url" if url else "html")
    meta = trafilatura.extract_metadata(html, default_url=url)
    if meta is not None:
        doc.title = doc.title or (meta.title or "")
        doc.author = meta.author or ""
    return doc


def fetch_url(url: str) -> ParsedDocument:
    import trafilatura

    if not re.match(r"^https?://", url):
        raise ValueError("Enter a full http(s) URL.")
    html = trafilatura.fetch_url(url)
    if not html:
        raise ValueError("Could not download that page. Check the URL or save the page and upload the file instead.")
    return clean_document(parse_html(html, url=url))
