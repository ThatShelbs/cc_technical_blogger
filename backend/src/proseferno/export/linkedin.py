"""LinkedIn: a feed post that sells the article, and a document carousel (PDF + editable PPTX).

LinkedIn's feed post limit is 3,000 characters and only the first ~210 show before
"…see more", so the hook comes first. Document posts accept PDF or PPTX; 1080x1350
(4:5 portrait) slides fill the most feed space on phones.
"""

from __future__ import annotations

import io
import re
import textwrap
from dataclasses import dataclass, field
from pathlib import Path

from proseferno.export.render import Figure
from proseferno.images import PLACEHOLDER_RE
from proseferno.ingest.stylometrics import split_sentences
from proseferno.writer import split_title

POST_LIMIT = 3000
SEE_MORE_AT = 210
SLIDE_W, SLIDE_H = 1080, 1350
MAX_SLIDES = 14

# Inferno palette.
INK, PLUM, MAGENTA, FLAME, AMBER, CREAM = "#1b0c41", "#4a0c6b", "#a52c60", "#ed6925", "#fb9b06", "#fcffa4"
PAPER = "#fff8e6"


@dataclass
class Section:
    heading: str
    paragraphs: list[str] = field(default_factory=list)
    slots: list[str] = field(default_factory=list)


@dataclass
class Slide:
    kind: str  # cover | section | image | closing
    title: str
    lines: list[str] = field(default_factory=list)
    image: Path | None = None
    caption: str = ""


# -- reading the post ------------------------------------------------------------------


def _clean_inline(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\*\*(.+?)\*\*|__(.+?)__", lambda m: m.group(1) or m.group(2), text)
    text = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"\1", text)
    return re.sub(r"`([^`]+)`", r"\1", text).strip()


def sections(markdown: str) -> tuple[list[str], list[Section]]:
    """Intro paragraphs and ``##`` sections (code blocks and lists are skipped)."""
    text = re.sub(r"```.*?```", "", markdown, flags=re.S)
    intro: list[str] = []
    out: list[Section] = []
    for block in re.split(r"\n\s*\n", text):
        block = block.strip()
        if not block:
            continue
        slot = PLACEHOLDER_RE.fullmatch(block)
        if slot:
            if out:
                out[-1].slots.append(slot.group(1))
            continue
        if block.startswith("## "):
            out.append(Section(heading=_clean_inline(block[3:].splitlines()[0])))
            continue
        if block.startswith("#") or block.startswith(("-", "*", "1.", ">", "|")):
            continue
        paragraph = _clean_inline(" ".join(block.splitlines()))
        (out[-1].paragraphs if out else intro).append(paragraph)
    return intro, out


def _first_sentences(text: str, count: int, limit: int) -> str:
    return textwrap.shorten(" ".join(split_sentences(text)[:count]), limit, placeholder="…")


# -- feed post --------------------------------------------------------------------------


def hashtags(tags: list[str], limit: int = 5) -> list[str]:
    out = []
    for tag in tags:
        words = re.findall(r"[A-Za-z0-9]+", tag)
        if words:
            out.append("#" + "".join(w if any(c.isupper() for c in w) else w.capitalize() for w in words))
    return list(dict.fromkeys(out))[:limit]


def template_post(markdown: str, tags: list[str]) -> str:
    """A ready-to-edit feed post built from the article itself (no model call)."""
    title, subtitle = split_title(markdown)
    intro, secs = sections(markdown)
    hook = _first_sentences(intro[0], 2, SEE_MORE_AT) if intro else subtitle or title
    lines = [hook, ""]
    if subtitle and subtitle not in hook:
        lines += [subtitle, ""]
    if secs:
        lines.append(f"In “{title}”:" if title else "In the article:")
        lines += [f"→ {s.heading}" for s in secs[:6]]
        lines.append("")
    lines += ["Full article in the comments (or on my profile).", "", "What would you add?"]
    if tags:
        lines += ["", " ".join(hashtags(tags))]
    return fit_post("\n".join(lines))


def fit_post(text: str) -> str:
    """Strip Markdown LinkedIn would show literally, and stay under the character limit."""
    text = re.sub(r"^#{1,6}\s+", "", text, flags=re.M)  # headings, not hashtags
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) <= POST_LIMIT:
        return text
    cut = text[: POST_LIMIT - 1]
    return cut[: cut.rfind("\n") if "\n" in cut[-400:] else len(cut)].rstrip() + "…"


# -- carousel ---------------------------------------------------------------------------


def carousel_slides(markdown: str, figures: dict[str, Figure], author: str) -> list[Slide]:
    title, subtitle = split_title(markdown)
    intro, secs = sections(markdown)

    def image(slot: str) -> Figure | None:
        fig = figures.get(slot)
        return fig if fig is not None and fig.path is not None and fig.path.exists() else None

    hero = image("hero")
    slides = [Slide("cover", title or "Untitled", [subtitle] if subtitle else [], hero.path if hero else None)]
    for section in secs:
        if len(slides) >= MAX_SLIDES - 1:
            break
        points = [_first_sentences(p, 1, 150) for p in section.paragraphs[:3]]
        slides.append(Slide("section", section.heading, points or ([_first_sentences(intro[0], 1, 150)] if intro else [])))
        for slot in section.slots:
            fig = image(slot)
            if fig is not None and slot not in ("hero", "meme") and len(slides) < MAX_SLIDES - 1:
                slides.append(Slide("image", section.heading, [], fig.path, fig.full_caption))
                break
    meme = image("meme")
    slides.append(
        Slide(
            "closing",
            "Read the full article",
            [title, f"by {author}" if author else "", "Follow for more like this."],
            meme.path if meme else None,
        )
    )
    return slides


def _rgb(hex_color: str) -> tuple[float, float, float]:
    value = hex_color.lstrip("#")
    return tuple(int(value[i : i + 2], 16) / 255 for i in (0, 2, 4))  # type: ignore[return-value]


def _latin1(text: str) -> str:
    """The PDF's built-in fonts cover Latin-1/WinAnsi; swap the few characters they lack."""
    return text.replace("→", "->").encode("cp1252", "replace").decode("cp1252")


def _jpeg(path: Path, max_width: int = 1000) -> bytes:
    """Slides show images at most ~900px wide; a resized JPEG keeps the deck small enough to upload."""
    from PIL import Image

    with Image.open(path) as source:
        img = source.convert("RGB")
    img.thumbnail((max_width, max_width * 2))
    out = io.BytesIO()
    img.save(out, format="JPEG", quality=88, optimize=True)
    return out.getvalue()


def _section_numbers(slides: list[Slide]) -> dict[int, int]:
    numbers: dict[int, int] = {}
    for i, slide in enumerate(slides):
        if slide.kind == "section":
            numbers[i] = len(numbers) + 1
    return numbers


def carousel_pdf(slides: list[Slide], title: str) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    total = len(slides)
    margin = 90
    numbers = _section_numbers(slides)

    def text_box(
        page: pymupdf.Page, rect: pymupdf.Rect, text: str, size: float, color: str, *, bold: bool = False, align: int = 0
    ) -> float:
        """Draw text, shrinking it until it fits; return the height it used."""
        text = _latin1(text)
        font = "hebo" if bold else "helv"
        while size > 18:
            # insert_textbox returns the unused height, or a negative number when the text does not fit.
            shape = page.new_shape()
            spare = shape.insert_textbox(rect, text, fontsize=size, fontname=font, color=_rgb(color), align=align)
            if spare >= 0:
                shape.commit()
                return rect.height - spare
            size -= 4
        page.insert_textbox(rect, text, fontsize=size, fontname=font, color=_rgb(color), align=align)
        return rect.height

    def image(page: pymupdf.Page, rect: pymupdf.Rect, path: Path) -> None:
        page.insert_image(rect, stream=_jpeg(path), keep_proportion=True)

    for index, slide in enumerate(slides, start=1):
        page = doc.new_page(width=SLIDE_W, height=SLIDE_H)
        dark = slide.kind in ("cover", "closing")
        page.draw_rect(page.rect, color=None, fill=_rgb(INK if dark else PAPER))
        page.draw_rect(pymupdf.Rect(0, 0, SLIDE_W, 18), color=None, fill=_rgb(FLAME))
        fg, accent = (CREAM, AMBER) if dark else (INK, MAGENTA)
        if slide.kind == "cover":
            if slide.image:
                image(page, pymupdf.Rect(margin, 110, SLIDE_W - margin, 640), slide.image)
            y = 690 + text_box(page, pymupdf.Rect(margin, 690, SLIDE_W - margin, 1010), slide.title, 76, fg, bold=True) + 40
            if slide.lines:
                text_box(page, pymupdf.Rect(margin, y, SLIDE_W - margin, 1190), slide.lines[0], 38, accent)
            text_box(page, pymupdf.Rect(margin, 1220, SLIDE_W - margin, 1290), "Swipe  >", 32, AMBER, bold=True, align=2)
        elif slide.kind == "section":
            text_box(page, pymupdf.Rect(margin, 110, SLIDE_W - margin, 160), f"{numbers[index - 1]:02d}", 34, FLAME, bold=True)
            y = 170 + text_box(page, pymupdf.Rect(margin, 170, SLIDE_W - margin, 470), slide.title, 68, fg, bold=True) + 70
            page.draw_rect(pymupdf.Rect(margin, y - 35, margin + 120, y - 29), color=None, fill=_rgb(AMBER))
            for line in slide.lines:
                page.draw_rect(pymupdf.Rect(margin, y + 12, margin + 14, y + 26), color=None, fill=_rgb(FLAME))
                y += text_box(page, pymupdf.Rect(margin + 40, y, SLIDE_W - margin, min(y + 260, SLIDE_H - 120)), line, 42, fg) + 50
        elif slide.kind == "image":
            text_box(page, pymupdf.Rect(margin, 100, SLIDE_W - margin, 260), slide.title, 52, fg, bold=True)
            if slide.image:
                image(page, pymupdf.Rect(margin, 290, SLIDE_W - margin, 1110), slide.image)
            if slide.caption:
                text_box(page, pymupdf.Rect(margin, 1130, SLIDE_W - margin, 1270), slide.caption, 26, MAGENTA)
        else:
            text_box(page, pymupdf.Rect(margin, 120, SLIDE_W - margin, 260), slide.title, 72, fg, bold=True)
            body = "\n".join(line for line in slide.lines if line)
            text_box(page, pymupdf.Rect(margin, 290, SLIDE_W - margin, 600), body, 44, accent)
            if slide.image:
                image(page, pymupdf.Rect(margin + 120, 640, SLIDE_W - margin - 120, 1240), slide.image)
        text_box(page, pymupdf.Rect(margin, SLIDE_H - 60, SLIDE_W - margin, SLIDE_H - 20), f"{index}/{total}", 22, accent, align=2)
    doc.set_metadata({"title": title, "creator": "Proseferno"})
    return doc.tobytes()


def carousel_pptx(slides: list[Slide]) -> bytes:
    from PIL import Image
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt

    px = 9525  # EMU per pixel at 96 dpi
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(SLIDE_W * px), Emu(SLIDE_H * px)
    blank = prs.slide_layouts[6]
    margin = 90

    def color(hex_color: str) -> RGBColor:
        return RGBColor.from_string(hex_color.lstrip("#").upper())

    def box(slide, x: int, y: int, w: int, h: int, text: str, size: int, fg: str, *, bold: bool = False):  # type: ignore[no-untyped-def]
        shape = slide.shapes.add_textbox(Emu(x * px), Emu(y * px), Emu(w * px), Emu(h * px))
        frame = shape.text_frame
        frame.word_wrap = True
        for i, line in enumerate(text.split("\n")):
            para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            run = para.add_run()
            run.text = line
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color(fg)
            run.font.name = "Segoe UI"

    def picture(slide, path: Path, x: int, y: int, w: int, h: int) -> None:  # type: ignore[no-untyped-def]
        with Image.open(path) as img:
            scale = min(w / img.width, h / img.height)
            pw, ph = int(img.width * scale), int(img.height * scale)
        slide.shapes.add_picture(
            io.BytesIO(_jpeg(path)), Emu((x + (w - pw) // 2) * px), Emu((y + (h - ph) // 2) * px), Emu(pw * px), Emu(ph * px)
        )

    total = len(slides)
    numbers = _section_numbers(slides)
    for index, data in enumerate(slides, start=1):
        slide = prs.slides.add_slide(blank)
        dark = data.kind in ("cover", "closing")
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = color(INK if dark else PAPER)
        fg, accent = (CREAM, AMBER) if dark else (INK, MAGENTA)
        width = SLIDE_W - 2 * margin
        if data.kind == "cover":
            if data.image:
                picture(slide, data.image, margin, 110, width, 530)
            box(slide, margin, 680, width, 330, data.title, 44, fg, bold=True)
            if data.lines:
                box(slide, margin, 1020, width, 160, data.lines[0], 22, accent)
        elif data.kind == "section":
            box(slide, margin, 100, width, 60, f"{numbers[index - 1]:02d}", 20, FLAME, bold=True)
            box(slide, margin, 170, width, 300, data.title, 40, fg, bold=True)
            box(slide, margin, 520, width, 700, "\n\n".join(f"• {line}" for line in data.lines), 24, fg)
        elif data.kind == "image":
            box(slide, margin, 100, width, 160, data.title, 30, fg, bold=True)
            if data.image:
                picture(slide, data.image, margin, 290, width, 820)
            if data.caption:
                box(slide, margin, 1130, width, 140, data.caption, 14, MAGENTA)
        else:
            box(slide, margin, 120, width, 140, data.title, 42, fg, bold=True)
            box(slide, margin, 290, width, 300, "\n".join(line for line in data.lines if line), 26, accent)
            if data.image:
                picture(slide, data.image, margin + 120, 640, width - 240, 600)
        box(slide, margin, SLIDE_H - 70, width, 50, f"{index}/{total}", 12, accent)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
